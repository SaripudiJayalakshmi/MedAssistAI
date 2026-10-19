from fpdf import FPDF

def generate_chat_pdf(question: str, answer: str, sources: list[str]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "MedAssist AI - Chat Export", ln=True)
    pdf.ln(5)

    pdf.set_font("Helvetica", "B", 12)
    pdf.multi_cell(0, 8, f"Question: {question}")
    pdf.ln(3)

    pdf.set_font("Helvetica", "", 11)
    clean_answer = answer.encode("latin-1", "ignore").decode("latin-1")
    pdf.multi_cell(0, 7, f"Answer:\n{clean_answer}")
    pdf.ln(3)

    if sources:
        pdf.set_font("Helvetica", "I", 10)
        pdf.multi_cell(0, 6, f"Sources: {', '.join(sources)}")

    return bytes(pdf.output())
from docx import Document
from io import BytesIO

def generate_chat_docx(question: str, answer: str, sources: list[str]) -> bytes:
    doc = Document()
    doc.add_heading("MedAssist AI - Chat Export", level=1)
    doc.add_heading("Question", level=2)
    doc.add_paragraph(question)
    doc.add_heading("Answer", level=2)
    for line in answer.split("\n"):
        line = line.strip().replace("**", "")
        if not line:
            continue
        if line.startswith("#"):
            doc.add_heading(line.lstrip("#").strip(), level=3)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    if sources:
        doc.add_heading("Sources", level=2)
        doc.add_paragraph(", ".join(sources))
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()